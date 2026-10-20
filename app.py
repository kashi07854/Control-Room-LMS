"""Control Room - Final Premium Build"""
import os, sys, json, docx, schedule, threading, time
import sqlite3, PyPDF2, docx, openpyxl, threading, time, urllib.parse, schedule
from docx.shared import Pt
import openpyxl
from pathlib import Path
from datetime import datetime, timedelta
from flask import Flask, render_template, request, jsonify

BASE = Path(__file__).resolve().parent
DB_PATH = BASE / "data" / "index.db"
DB_PATH.parent.mkdir(parents=True, exist_ok=True)

def db_conn():
    c = sqlite3.connect(DB_PATH, timeout=30)
    c.row_factory = sqlite3.Row
    return c

def init_db():
    c = db_conn()
    c.execute("CREATE TABLE IF NOT EXISTS items (id INTEGER PRIMARY KEY AUTOINCREMENT, parent_id INTEGER, name TEXT, rel_path TEXT UNIQUE, item_type TEXT, ext TEXT, size INTEGER, is_deleted INTEGER DEFAULT 0)")
    c.commit(); c.close()

SCAN_STATE = {"running": False, "done": 0}
def scan_worker():
    SCAN_STATE["running"] = True
    c = db_conn()
    root = Path("G:\\My Drive")
    try:
        for path in root.rglob("*"):
            if "00_CONTROL_ROOM" in str(path): continue
            rel = str(path.relative_to(root)).replace("\\", "/")
            is_dir = path.is_dir()
            ext = "" if is_dir else path.suffix.lower().lstrip(".")
            
            existing = c.execute("SELECT id FROM items WHERE rel_path=?", (rel,)).fetchone()
            if not existing:
                c.execute("INSERT INTO items (name, rel_path, item_type, ext, size) VALUES (?,?,?,?,?)",
                          (path.name, rel, "folder" if is_dir else "file", ext, 0))
                SCAN_STATE["done"] += 1
        c.commit()
    except Exception as e:
        print(f"Scan error: {e}")
    finally:
        c.close()
        SCAN_STATE["running"] = False

app = Flask(__name__)

@app.route("/")
def dashboard():
    return render_template("dashboard.html")

@app.route("/browse")
def browse():
    rel = request.args.get("path", "").replace("\\", "/").strip("/")
    c = db_conn()
    try:
        if rel:
            parent = c.execute("SELECT id FROM items WHERE rel_path=?", (rel,)).fetchone()
            items = c.execute("SELECT * FROM items WHERE parent_id=? OR rel_path LIKE ?", (parent["id"] if parent else -1, rel + "/%")).fetchall()
            # Filter only direct children for simplicity
            direct = [i for i in items if i["rel_path"].rsplit("/", 1)[0] == rel] if parent else items
        else:
            direct = c.execute("SELECT * FROM items WHERE parent_id IS NULL").fetchall()
        breadcrumbs = [{"name": p, "path": "/".join(rel.split("/")[:i+1])} for i, p in enumerate(rel.split("/"))] if rel else []
        return render_template("browse.html", items=direct, rel=rel, breadcrumbs=breadcrumbs, error=None)
    finally: c.close()

@app.route("/api/save_life_os", methods=["POST"])
def save_life_os():
    data = request.json
    excel_path = BASE / "Life_OS_SIMPLE.xlsx"
    if not excel_path.exists(): excel_path = BASE / "data" / "Life_OS_SIMPLE.xlsx"
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb[data["sheet"]]
        ws.cell(row=data["row"] + 1, column=data["col"] + 1).value = data["value"]
        wb.save(excel_path)
        return jsonify({"ok": True})
    except Exception as e: return jsonify({"ok": False, "error": str(e)}), 500

@app.route("/folder_tree")
def folder_tree():
    excel_path = BASE / "ONLY_FOLDERS.xlsx"
    if not excel_path.exists(): excel_path = BASE / "data" / "ONLY_FOLDERS.xlsx"
    if not excel_path.exists(): return render_template("folder_tree.html", error="ONLY_FOLDERS.xlsx not found.")
    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        folders = []
        for row in ws.iter_rows(min_row=2, values_only=False):
            category, cell, level = row[0].value, row[1], row[2].value
            hyperlink = cell.hyperlink.target if cell.hyperlink else None
            display_text = cell.value
            if not hyperlink and isinstance(display_text, str) and "HYPERLINK" in display_text:
                import re
                match = re.search(r'HYPERLINK\("([^"]+)"', display_text)
                if match: hyperlink = match.group(1)
            if hyperlink:
                rel_path = hyperlink.replace("G:\\My Drive\\", "").replace("\\", "/")
                web_link = f"/browse?path={urllib.parse.quote(rel_path)}"
                folders.append({"category": category, "name": display_text, "level": level, "link": web_link})
        return render_template("folder_tree.html", folders=folders, error=None)
    except Exception as e: return render_template("folder_tree.html", error=f"Error: {e}")

@app.route("/api/scan", methods=["POST"])
def api_scan():
    if SCAN_STATE["running"]: return jsonify({"ok": False})
    threading.Thread(target=scan_worker, daemon=True).start()
    return jsonify({"ok": True})

@app.route("/api/scan_status")
def api_scan_status(): return jsonify(SCAN_STATE)

@app.route("/habit_tracker")
def habit_tracker():
    return render_template("habit_tracker.html")

@app.route("/habit_tracker_raw")
def habit_tracker_raw():
    return render_template("habit_tracker_raw.html")

@app.route("/api/open_file/<int:item_id>", methods=["POST"])
def open_file(item_id):
    c = db_conn()
    try:
        item = c.execute("SELECT abs_path FROM items WHERE id=?", (item_id,)).fetchone()
        if item and os.path.exists(item["abs_path"]):
            os.startfile(item["abs_path"]) # Opens with default Windows app
            return jsonify({"ok": True})
        return jsonify({"ok": False, "error": "File not found on disk"}), 404
    finally: c.close()

@app.route("/diary_tracker")
def diary_tracker():
    return render_template("diary_tracker.html")

@app.route("/diary_tracker_raw")
def diary_tracker_raw():
    return render_template("diary_tracker_raw.html")

@app.route("/chronodeck")
def chronodeck():
    return render_template("chronodeck.html")

@app.route("/chronodeck_raw")
def chronodeck_raw():
    return render_template("chronodeck_raw.html")

@app.route("/ai")
def ai_chat():
    return render_template("ai_chat.html")

@app.route("/api/ai/chat", methods=["POST"])
def api_ai_chat():
    import urllib.request
    import json
    
    data = request.get_json(silent=True) or {}
    prompt = data.get("prompt")
    model = data.get("model", "llama3")
    
    if not prompt:
        return jsonify({"ok": False, "error": "Prompt empty"}), 400

    try:
        req = urllib.request.Request(
            "http://127.0.0.1:11434/api/generate",
            data=json.dumps({"model": model, "prompt": prompt, "stream": False}).encode('utf-8'),
            headers={'Content-Type': 'application/json'}
        )
        with urllib.request.urlopen(req, timeout=90) as response:
            res = json.loads(response.read().decode())
            return jsonify({"ok": True, "response": res.get("response", "")})
    except urllib.error.URLError:
        return jsonify({"ok": False, "error": "Ollama connect nahi hua. Kya aapne Ollama install aur run kiya hai? ('ollama run llama3')"}), 500
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

def extract_text_from_file(file_path):
    ext = file_path.lower().split('.')[-1]
    text = ""
    try:
        if ext == 'pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages[:10]: # Limit to first 10 pages to save context
                    text += (page.extract_text() or "") + "\n"
        elif ext in ['docx', 'doc']:
            doc = docx.Document(file_path)
            for para in doc.paragraphs[:200]: # Limit paragraphs
                text += para.text + "\n"
        elif ext in ['xlsx', 'xls', 'csv']:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(max_row=50, values_only=True):
                text += " | ".join([str(c) for c in row if c]) + "\n"
        elif ext in ['txt', 'md', 'py', 'json', 'html', 'css', 'js']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
    except Exception as e:
        text = f"Error reading file: {e}"
    return text[:15000] # Send max ~15,000 characters to fit AI memory

@app.route("/api/all_files")
def api_all_files():
    c = db_conn()
    try:
        # Only fetch valid document files and ignore MacOS junk (.DS_Store, ._*, etc)
        files = c.execute("""
            SELECT id, name, rel_path FROM items 
            WHERE item_type='file' 
            AND is_deleted=0 
            AND name NOT LIKE '.%' 
            AND name NOT LIKE '~$%'
            AND ext IN ('pdf', 'docx', 'doc', 'xlsx', 'xls', 'txt', 'md', 'csv')
            ORDER BY name LIMIT 300
        """).fetchall()
        return jsonify([{"id": f["id"], "name": f["name"], "path": f["rel_path"]} for f in files])
    finally: c.close()
@app.route("/api/ai/chat_file", methods=["POST"])
def api_ai_chat_file():
    import urllib.request
    import json
    
    data = request.get_json(silent=True) or {}
    prompt = data.get("prompt")
    file_id = data.get("file_id")
    model = data.get("model", "llama3")
    
    if not prompt or not file_id:
        return jsonify({"ok": False, "error": "Prompt or File missing"}), 400

    c = db_conn()
    try:
        item = c.execute("SELECT abs_path, name FROM items WHERE id=?", (file_id,)).fetchone()
    finally: c.close()

    if not item or not os.path.exists(item["abs_path"]):
        return jsonify({"ok": False, "error": "File not found on disk."}), 404

    # Extract text from the local file
    file_text = extract_text_from_file(item["abs_path"])
    
    if not file_text.strip():
        return jsonify({"ok": False, "error": "Could not extract text from this file. It might be an image-based PDF."}), 400

    # Create a prompt containing the file context
    full_prompt = f"""
    You are a helpful assistant. I am providing you the contents of a file named '{item['name']}'.
    Please read the file content below and answer my question based ONLY on this content.
    
    --- FILE CONTENT START ---
    {file_text}
    --- FILE CONTENT END ---
    
    MY QUESTION: {prompt}
    """

    try:
        req = urllib.request.Request(
            "http://127.0.0.1:11434/api/generate",
            data=json.dumps({"model": model, "prompt": full_prompt, "stream": False}).encode('utf-8'),
            headers={'Content-Type': 'application/json'}
        )
        with urllib.request.urlopen(req, timeout=120) as response:
            res = json.loads(response.read().decode())
            return jsonify({"ok": True, "response": res.get("response", "")})
    except urllib.error.URLError:
        return jsonify({"ok": False, "error": "Ollama connect nahi hua. Is 'ollama serve' running?"}), 500
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

# ================= REPORT GENERATION & AUTO-SYNC =================
import docx
from docx.shared import Pt

REPORTS_DIR = BASE / "Reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

def create_excel_report(data, report_type):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{report_type} Report"
    
    # Styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    headers = ["Date", "Morning", "Accomplished", "Learned", "Evening", "Notes"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    row = 2
    entries = data.get("journal", {}).get("entries", {})
    for date, entry in sorted(entries.items()):
        ws.cell(row=row, column=1, value=date)
        ws.cell(row=row, column=2, value=entry.get("morning", ""))
        ws.cell(row=row, column=3, value=entry.get("acc", ""))
        ws.cell(row=row, column=4, value=entry.get("learn", ""))
        ws.cell(row=row, column=5, value=entry.get("evening", ""))
        ws.cell(row=row, column=6, value=entry.get("notes", ""))
        row += 1
        
    # Auto-adjust width
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
        
    filename = REPORTS_DIR / f"{report_type}_Report_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    wb.save(filename)
    return str(filename)

def create_word_report(data, report_type):
    doc = docx.Document()
    title = doc.add_heading(f'MYLIFE OS - {report_type} Report', 0)
    
    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%d %b %Y, %I:%M %p')}")
    doc.add_paragraph("-" * 50)
    
    entries = data.get("journal", {}).get("entries", {})
    for date, entry in sorted(entries.items()):
        doc.add_heading(f"Date: {date}", level=1)
        if entry.get("morning"): doc.add_paragraph(f"🌅 Morning: {entry['morning']}")
        if entry.get("acc"): doc.add_paragraph(f"🎯 Accomplished: {entry['acc']}")
        if entry.get("learn"): doc.add_paragraph(f"📚 Learned: {entry['learn']}")
        if entry.get("evening"): doc.add_paragraph(f"🌙 Evening: {entry['evening']}")
        if entry.get("notes"): doc.add_paragraph(f"💡 Notes: {entry['notes']}")
        doc.add_paragraph("-" * 20)
        
    filename = REPORTS_DIR / f"{report_type}_Report_{datetime.now().strftime('%Y-%m-%d')}.docx"
    doc.save(filename)
    return str(filename)

@app.route("/api/sync_diary", methods=["POST"])
def sync_diary():
    data = request.get_json()
    if not data: return jsonify({"ok": False}), 400
    try:
        with open(BASE / "data" / "diary_sync.json", "w", encoding="utf-8") as f:
            json.dump(data, f)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app.route("/api/generate_report", methods=["POST"])
def generate_report():
    data = request.get_json()
    if not data: return jsonify({"ok": False, "error": "No data received"}), 400
    rtype = "Weekly" if datetime.now().weekday() == 6 else "Monthly" # Sunday = 6, adjust logic as needed
    # Force type based on frontend request
    rtype = data.get("type", "Auto")
    try:
        xlsx_path = create_excel_report(data, rtype)
        docx_path = create_word_report(data, rtype)
        return jsonify({"ok": True, "excel": xlsx_path, "word": docx_path})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

def scheduled_report_checker():
    import time
    while True:
        now = datetime.now()
        # Check if it's Sunday at 23:59 (End of week)
        if now.weekday() == 6 and now.hour == 23 and now.minute >= 59:
            sync_file = BASE / "data" / "diary_sync.json"
            if sync_file.exists():
                try:
                    with open(sync_file, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    create_excel_report(data, "Weekly_Auto")
                    create_word_report(data, "Weekly_Auto")
                except: pass
            time.sleep(120) # Avoid duplicate runs in same minute
        # Check if it's last day of month at 23:59
        elif (now + timedelta(days=1)).month != now.month and now.hour == 23 and now.minute >= 59:
            sync_file = BASE / "data" / "diary_sync.json"
            if sync_file.exists():
                try:
                    with open(sync_file, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    create_excel_report(data, "Monthly_Auto")
                    create_word_report(data, "Monthly_Auto")
                except: pass
            time.sleep(120)
        time.sleep(60)

if __name__ == "__main__":
    init_db()
    threading.Thread(target=scheduled_report_checker, daemon=True).start()
    app.run(host="127.0.0.1", port=5000, debug=False, use_reloader=False)



















